from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from research_forge.capability_registry import load_capability_registry
from research_forge.upstream_capabilities import (
    AdoptionStatus,
    UPSTREAM_CAPABILITIES,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Research_Forge_功能与GitHub能力汇总_2026-07-31_final.docx"

INK = "051D25"
DEEP = "0D454E"
MID = "3A747D"
REEF = "70A1A9"
PALE = "EAF3F4"
LIGHT = "F5F9F9"
GREEN = "168A4A"
AMBER = "9A6A00"
RED = "A43B2B"

STATUS_CN = {
    AdoptionStatus.VERIFIED_RUNTIME: "已验证运行",
    AdoptionStatus.PARTIAL_RUNTIME: "部分运行",
    AdoptionStatus.ADAPTER_ONLY: "产物适配",
    AdoptionStatus.ARCHITECTURE_ONLY: "架构借鉴",
    AdoptionStatus.CODEX_SKILL_ONLY: "仅 Codex Skill",
    AdoptionStatus.STUB_OR_MISSING: "未实现/缺失",
    AdoptionStatus.NOT_ADOPTED: "未采用",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_font(run, size=10.5, bold=False, color=INK, italic=False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size={1: 17, 2: 13.5, 3: 11.5}[level], bold=True, color=DEEP)


def add_body(doc: Document, text: str, *, color=INK, italic=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    add_text(p, text, size=10.5, color=color, italic=italic)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, size=10.2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], DEEP)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_text(p, header, size=9.3, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT)
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=8.7, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc: Document, label: str, text: str, fill=PALE, color=DEEP) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.18
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), REEF)
        node.set(qn("w:space"), "5")
        borders.append(node)
    p_pr.append(borders)
    add_text(p, label + "  ", size=10.2, bold=True, color=color)
    add_text(p, text, size=10.2, color=INK)


def status_summary() -> list[tuple[str, int]]:
    return [
        (STATUS_CN[status], sum(item.status is status for item in UPSTREAM_CAPABILITIES))
        for status in AdoptionStatus
    ]


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for level in (1, 2, 3):
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string(DEEP)
        style.paragraph_format.space_before = Pt({1: 15, 2: 11, 3: 8}[level])
        style.paragraph_format.space_after = Pt({1: 7, 2: 5, 3: 3}[level])

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "RESEARCH FORGE  ·  CAPABILITY AUDIT", size=8.5, bold=True, color=MID)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Evidence-gated capability register · 2026-07-31", size=8, color=REEF)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(7)
    add_text(p, "能力审计", size=10, bold=True, color=MID)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "Research Forge", size=30, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    add_text(p, "现有功能与 GitHub / Skill 能力汇总", size=18, color=DEEP)
    add_callout(
        doc,
        "结论",
        "平台已经具备四阶段控制闭环、检索与证据绑定、窄范围实验执行、论文与审计能力；"
        "但尚未达到任意课题的跨领域自主科研闭环或外部独立复现 C5。",
    )
    add_body(doc, "版本日期：2026-07-31　｜　证据口径：代码 + 测试 + 持久真实运行记录", color=MID)
    add_body(doc, "本文把“参考过某个 GitHub”与“平台已经真实运行其能力”严格分开。", italic=True)

    doc.add_page_break()
    add_heading(doc, "1. 读表口径", 1)
    maturity_rows = [
        ["C0", "概念", "只有目标或设计"],
        ["C1", "已实现", "有源代码，测试仍不足"],
        ["C2", "组件验证", "代码和组件测试通过"],
        ["C3", "受控端到端", "在受控环境中完整走通"],
        ["C4", "真实案例验证", "具有可定位、可重验的真实记录"],
        ["C5", "独立验证", "隔离信任域或外部团队独立复现"],
    ]
    add_table(doc, ["等级", "名称", "允许说明"], maturity_rows, [900, 1700, 6760])
    add_callout(
        doc,
        "科学权限边界",
        "外部检索、NLI、AI 科学家面板和论文 Agent 均不能改写冻结的科学 Verdict；"
        "执行失败也不等于假设被反驳。",
        fill="FFF5DA",
        color=AMBER,
    )

    add_heading(doc, "2. 四阶段产品流程", 1)
    phase_rows = [
        ["01", "方向发现", "只读扫描项目、抽取 claim / 概念、检索外部信号、形成候选组合并冻结 Scope"],
        ["02", "协议与可行性", "收敛问题与假设、确定指标和数据边界、完成 MVP、形成并批准 Research Contract"],
        ["03", "实验与判定", "构建正式资产、隔离运行、独立评价、证据绑定、统计分析与 Verdict"],
        ["04", "论文与审计", "证据约束写作、图表、审计、AI 科学评审、作者投稿批准；证据不足时回退补证据"],
    ]
    add_table(doc, ["阶段", "名称", "主要交付"], phase_rows, [720, 1700, 6940])

    add_heading(doc, "3. 已审计平台能力", 1)
    capabilities = load_capability_registry()
    cap_rows = []
    for item in capabilities:
        cap_rows.append([
            f"C{int(item.declared_maturity)}",
            item.title,
            item.claim,
            item.claim_boundary,
        ])
    add_table(doc, ["成熟度", "能力", "已达到", "边界"], cap_rows, [820, 1900, 3300, 3340])

    doc.add_page_break()
    add_heading(doc, "4. GitHub / Skill 接入总览", 1)
    summary_rows = [[name, str(count)] for name, count in status_summary()]
    add_table(doc, ["接入状态", "数量"], summary_rows, [6500, 2860])
    add_body(
        doc,
        "注册表共 42 项；机器审计为 42/42 声明均与当前证据等级一致。这里的“通过”表示没有越级宣传，"
        "不表示 42 个上游项目全部已经接入。",
    )

    for status in AdoptionStatus:
        items = [item for item in UPSTREAM_CAPABILITIES if item.status is status]
        if not items:
            continue
        add_heading(doc, STATUS_CN[status], 2)
        rows = []
        for item in items:
            landed = "；".join(item.learned_capabilities)
            boundary = "；".join(item.missing_or_bounded)
            rows.append([
                item.name,
                landed,
                item.permitted_claim,
                boundary,
            ])
        add_table(doc, ["来源", "RF 已吸收/实现", "当前允许声称", "仍有边界"], rows, [1600, 2700, 2500, 2560])

    add_heading(doc, "5. 当前能力判断", 1)
    add_heading(doc, "已经达到平台可用或更高", 2)
    for text in (
        "双入口、四阶段 Study / StepInstance / DAG / Gate 控制闭环",
        "外部检索网关的离线默认、授权、净化、预算、幂等与冻结响应",
        "PaperQA 冻结全文证据分析；draw.io 可编辑源文件与 Desktop 导出",
        "OpenML 五任务窄范围受控 Benchmark",
        "AIRS 四个官方 RAD 任务、40 个隔离运行单元的窄范围 C4",
        "两个真实 Python 项目的同主机密封双工作区重放",
        "基础 RO-Crate 1.1 已通过维护中的外部 checker 离线验收",
        "Great Expectations 冻结表格结构质量门已完成离线受控验收",
        "受限 Stage 3 W3C PROV 投影已通过外部 PySHACL 语义引擎",
        "两份密封包已完成断网、只读、全新 Docker worker 的受控隔离复现",
        "证据约束论文流水线、故障定位 / Repair / Successor 与 Publication Readiness",
    ):
        add_bullet(doc, text)

    add_heading(doc, "仍为部分能力", 2)
    for text in (
        "任意跨领域 Research Contract 的通用编译和实验 Profile 覆盖",
        "Nuwa 面板真实模型矩阵与外部科学校准",
        "next-ai-draw-io 实时自然语言编辑闭环",
        "完整 W3C PROV-CONSTRAINTS 与 RO-Crate Workflow Run / Provenance Run Profile",
        "完整 AIRS 矩阵、排行榜验收与 AIRA Dojo",
        "外部组织或独立云账户 + KMS/HSM 回执的 clean-room 复现 C5",
    ):
        add_bullet(doc, text)

    add_heading(doc, "优先补齐顺序", 2)
    for text in (
        "把已完成的基础 RO-Crate 与两次本机隔离 replay 提升到 Workflow Run Profile 和外部信任域验收",
        "扩展真实 Stage 2 → Stage 3 科学链，而非只扩展工作流步骤数量",
        "为 Nuwa、next-ai-draw-io 和尚未覆盖的数据类型补真实运行验收",
        "建立外部盲审或隔离信任域证据，谨慎从 C4 走向 C5",
    ):
        add_bullet(doc, text)

    add_callout(
        doc,
        "最终判断",
        "Research Forge 已经是一个可运行、可审计的研究工作台；其强项是边界收敛、流程治理、"
        "证据绑定、故障定位与受约束写作。当前最大缺口不是更多页面，而是更多实验类型的真实科学闭环和独立复现。",
        fill=PALE,
    )

    doc.core_properties.title = "Research Forge 现有功能与 GitHub / Skill 能力汇总"
    doc.core_properties.subject = "Evidence-gated capability audit"
    doc.core_properties.author = "Research Forge"
    doc.core_properties.keywords = "Research Forge, capability audit, GitHub, scientific workflow"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
