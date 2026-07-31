from __future__ import annotations

"""Build the human-readable Research Forge capability and upstream catalog.

The TXT uses UTF-8 with BOM so legacy Windows editors render Chinese correctly.
The DOCX uses a compact reference-guide layout and CJK-safe fonts.
"""

from collections import Counter
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from research_forge.upstream_capabilities import (  # noqa: E402
    UPSTREAM_CAPABILITIES,
    AdoptionStatus,
    audit_upstream_capabilities,
)
from research_forge.capability_registry import audit_capability_registry  # noqa: E402


DATE = "2026-08-01"
TXT_PATH = ROOT / "docs" / f"Research_Forge_现有功能与GitHub能力总表_{DATE}.txt"
DOCX_PATH = ROOT / "docs" / f"Research_Forge_现有功能与GitHub能力总表_{DATE}.docx"

# compact_reference_guide with a named Research Forge ocean-palette override.
INK = "051D25"
DEEP = "0D454E"
MID = "3A747D"
PALE = "DDEBED"
LIGHT = "EEF5F6"
MUTED = "58737A"
GREEN = "168A4A"
AMBER = "A35D00"
RED = "A53A2B"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"


STATUS_CN = {
    AdoptionStatus.VERIFIED_RUNTIME: "已实跑接入",
    AdoptionStatus.PARTIAL_RUNTIME: "部分运行",
    AdoptionStatus.ADAPTER_ONLY: "仅适配器",
    AdoptionStatus.ARCHITECTURE_ONLY: "仅架构参考",
    AdoptionStatus.CODEX_SKILL_ONLY: "仅 Codex Skill",
    AdoptionStatus.STUB_OR_MISSING: "尚未实现",
    AdoptionStatus.NOT_ADOPTED: "未采用",
}


CORE_CAPABILITIES = [
    ("统一四阶段 Study 工作流", "已达到（C3）", "双入口、Study 隔离、持久化 StepInstance/DAG、负责人 Gate、暂停/恢复/重试和阶段交接。", "证明流程控制闭环，不保证任意课题都能完成科学实验。"),
    ("本地项目只读扫描与 Claim 发现", "已达到", "扫描代码、数据、报告、Office 文档和 HDF5 元数据；排除缓存、驱动包与共享二进制目录；形成候选主张和学术化概念。", "语义发现链只有显式哈希/标识绑定后才能升级为正式证据。"),
    ("外部检索网关", "已达到（C3）", "默认离线；负责人批准后统一访问论文、GitHub、公开网页、开放全文和精确 OpenML 资源，并记录预算、幂等、查询摘要和响应哈希。", "Provider 安装不等于联网已获批；机构登录与正式发布仍需用户授权。"),
    ("PaperQA 全文证据分析", "已达到（C4，限定）", "固定版本、冻结 corpus/index、EvidenceSpan、引文与哈希绑定；已有真实运行验收。", "不宣称普遍优于 PaperQA；完整答案由 Research Forge 的证据约束合成层生成。"),
    ("阶段一 Discovery Portfolio", "已达到", "把本地主张、学术关键词、外部文献/趋势信号和候选贡献统一比较，由负责人选择并冻结 Scope。", "趋势与热度只能辅助选题，不能决定科学 Verdict。"),
    ("阶段二可行性、MVP 与 Research Contract", "部分达到（C2）", "定义问题、假设、基线/实验组、指标、分母、样本原则、资源需求和最小冒烟验证；Contract Compiler 阻止占位字段进入正式实验。", "三类 Profile 已覆盖，但通用跨领域契约编译仍未完成。"),
    ("类型化实验 Profile", "部分达到（C2）", "已实现 tabular_ml_v1、benchmark_prediction_v1、existing_python_project_v1，以及有限的回归、分类、检索、问卷与投资类扩展。", "未知实验类型会明确阻断，不会伪装成功。"),
    ("OpenML 隐藏标签配对实验", "已达到（C4）", "OpenML Task 39 的候选与隐藏标签隔离；基线/处理在隔离容器运行，独立 evaluator 与冻结参考实现一致。", "是一个真实窄案例，不代表完整 OpenML Run/Flow 兼容，也不是外部 C5 复现。"),
    ("官方 AIRS RAD 实验", "已达到（C4，限定）", "四个官方任务、10 个种子、40 个隔离单元使用未修改官方 evaluator 完成。", "不是完整 AIRS 矩阵，也未提交排行榜。"),
    ("既有 Python 项目密封复现", "已达到（C4，限定）", "两个非预制真实项目快照在全新离线 Docker 中重放，冻结指标与重放结果一致。", "仍属于本地控制的真实案例，不是外部独立信任域。"),
    ("Great Expectations 结构资格门", "已达到（C3）", "对冻结表格资源运行哈希绑定、离线结构检查，并归一化为资源资格记录。", "只能证明结构质量，不能证明标签语义、因果有效性或科学结论。"),
    ("独立评估、证据链与 Verdict", "部分达到", "Protocol→Run→Output→Evaluation→Claim 显式绑定；确定性检查优先，NLI 只做风险告警；执行失败与科学反驳分离。", "仍需更多 Profile 和第二评估器家族的真实比较。"),
    ("W3C PROV 交换", "已达到（C3，限定）", "Stage 3 的受限 PROV 投影通过外部 PySHACL 语义引擎和冻结 SHACL 规则。", "不宣称覆盖全部 PROV-CONSTRAINTS，也不是 W3C 认证。"),
    ("RO-Crate 交换", "已达到（C3，Workflow Run）", "Stage 3 Workflow Run RO-Crate 0.5 连同继承的 Process Run、Workflow RO-Crate 与基础 RO-Crate 必需检查全部通过外部 CRS4 validator。", "Provenance Run Profile 仍待验收；不是第三方认证。"),
    ("故障定位、Repair 与有界回滚", "已达到（受控）", "追加式诊断、最早可预防阶段、Repair Contract、Artifact Dependency Graph、successor run 和历史 Verdict 保留。", "AI 只能建议故障来源；科学修复和高成本重跑仍需负责人批准。"),
    ("受控 clean-room replay", "已达到（C3/RF-E1）", "两个密封包在全新、离线、只读根文件系统、非 root Docker worker 中重算并比较。", "没有外部 KMS/WORM 收据和独立操作者，因此不能称 C5 独立复现。"),
    ("公共最小复现包", "已公开发布（C4）", "OpenML Task 39 隐藏标签复现包和正确阻断契约案例已公开发布，并在 GitHub-hosted runner 重放；收据具有可验证的 Sigstore/SLSA provenance。", "工作流仍由同一仓库所有者控制，尚无独立科学操作者签署，因此不是 C5。"),
    ("证据约束论文写作", "已达到（C3）", "自然学术标题、分层大纲、单段摘要、章节写作、数字/引用/Claim/结构审计、证据不足披露和 PDF 编译。", "不能把缺失证据美化成结果，也不能保证特定期刊录用。"),
    ("sci-ssci 写作约束", "部分达到", "已本地适配 reader-question 章节规划、目标期刊建模和证据保持型润色。", "没有嵌入上游完整 Skill runtime/corpus builder。"),
    ("Nuwa AI 科学家面板", "部分达到（C2）", "盲化 Claim 包、Feynman/Tukey/Shannon/Popper 风格角色、固定路由、veto/abstain 聚合和追加裁决队列。", "缺真实模型端到端验收和专门的人类裁决 UI；不得冒充真人外审。"),
    ("可编辑科研图与导出", "已达到（C4）", "生成、审计 .drawio，并由本地 draw.io Desktop 导出 SVG/PDF；论文流水线保留可编辑源。", "next-ai-draw-io 的实时 MCP 对话编辑仍只有适配器，未完成最新端到端验收。"),
    ("Completion Record 与签名", "部分达到", "记录 Study/Contract/Run、关键哈希、readiness 三层状态和独立验证命令；支持自有 Ed25519 不可变记录。", "不是第三方认证、SLSA provenance 或 Sigstore 透明日志。"),
]


PHASES = [
    ("阶段一：方向发现", "只读项目扫描 → Claim/关键词 → 外部文献与趋势 → Discovery Portfolio → 负责人冻结 Scope。", "可用；外部检索受项目网络策略约束。"),
    ("阶段二：协议与可行性", "收敛议题、假设、概念基线/处理、指标/分母、资源包和 MVP；编译并批准 Research Contract。", "可用但仍是主要泛化短板；已支持三类正式 Profile。"),
    ("阶段三：实验与判定", "构建正式资产、隔离运行、独立评估、资格门、统计规则、Hypothesis/Study Verdict、复现与 provenance。", "多个真实窄案例达到 C4；跨领域通用实验仍未达到。"),
    ("阶段四：论文与审计", "证据约束写作、论文图表、科学面板、数字/引用/主张审计、补证据回退、PDF 和投稿批准。", "写作控制闭环可用；真人外审与期刊录用不属于自动完成项。"),
]


def compact(text: str, limit: int = 260) -> str:
    value = " ".join(str(text).split())
    return value if len(value) <= limit else value[: limit - 1].rstrip() + "…"


def build_txt() -> str:
    cap_audit = audit_capability_registry(ROOT)
    up_audit = audit_upstream_capabilities(ROOT)
    status_counts = Counter(item.status for item in UPSTREAM_CAPABILITIES)
    lines: list[str] = []
    add = lines.append
    add("Research Forge 现有功能与 GitHub / Skill 能力总表")
    add(f"更新日期：{DATE}")
    add("")
    add("一、结论摘要")
    add("=" * 72)
    add("Research Forge 已形成可运行的四阶段研究工作台，并在流程控制、检索、证据约束写作、科研绘图、部分实验 Profile 和复现检查上取得真实验收。")
    add("当前最强结论不是‘可以自动完成任意科研课题’，而是：在已支持的计算实验类型中，平台能够把方向、契约、执行、证据、判定、修复和论文保持在同一条可审计链上。")
    add(f"机器注册表审计：{cap_audit.supported_count}/{cap_audit.manifest_count} 项产品能力声明有本地证据；{up_audit.supported_count}/{up_audit.record_count} 项上游采用状态通过证据定位审计。")
    add("独立复现边界：已完成本地 clean-room replay，但没有外部独立操作者和外部签署收据，因此尚未达到 C5。")
    add("")
    add("成熟度：C0 概念；C1 实现；C2 组件验证；C3 受控端到端；C4 真实案例；C5 外部独立信任域复现。")
    add("")
    add("二、四阶段当前能力")
    add("=" * 72)
    for title, flow, state in PHASES:
        add(title)
        add(f"  过程：{flow}")
        add(f"  现状：{state}")
    add("")
    add("三、核心功能清单")
    add("=" * 72)
    for index, (name, status, ability, boundary) in enumerate(CORE_CAPABILITIES, 1):
        add(f"{index}. {name}｜{status}")
        add(f"   已有：{ability}")
        add(f"   边界：{boundary}")
    add("")
    add("四、GitHub / Skill / 标准来源总体分布")
    add("=" * 72)
    for status in AdoptionStatus:
        add(f"{STATUS_CN[status]}：{status_counts.get(status, 0)} 项")
    add("说明：‘仅架构参考’不等于已安装；‘仅 Codex Skill’不等于 Research Forge 产品运行时会自动调用；‘已实跑接入’也只允许声明注册表中的限定范围。")
    add("")
    add("五、逐项上游能力与接入情况")
    add("=" * 72)
    for status in AdoptionStatus:
        records = [item for item in UPSTREAM_CAPABILITIES if item.status is status]
        if not records:
            continue
        add(f"[{STATUS_CN[status]}｜{len(records)} 项]")
        for item in records:
            add(f"- {item.name}")
            add(f"  来源：{item.source_url}")
            add(f"  上游能力：{compact('；'.join(item.upstream_capabilities), 420)}")
            learned = "；".join(item.learned_capabilities) or "未转化为平台能力"
            add(f"  Research Forge 采用：{compact(learned, 340)}")
            add(f"  可声明：{item.permitted_claim}")
            add(f"  未达到/限制：{compact('；'.join(item.missing_or_bounded), 420)}")
        add("")
    add("六、最重要的未完成项")
    add("=" * 72)
    gaps = [
        "C5 外部独立复现：需要独立操作者/信任域仅使用冻结包复现，并生成可验证外部收据。",
        "通用 Contract Compiler：继续扩展实验 Profile、可执行指标、目标规则、采样框与资源授权，减少阶段三退回阶段二。",
        "Provenance Run RO-Crate：Workflow Run 0.5 已验收，Provenance Run Profile 仍待完成。",
        "Nuwa 面板：补真实模型端到端运行和负责人裁决 UI。",
        "next-ai-draw-io：补真实 MCP 对话编辑与导出验收。",
        "大规模多文献 PaperQA benchmark：补多文献冲突、攻击、弃权准确率、引用准确率与 P95/P99。",
        "尚未接入：OSF 在线注册、MLflow、DVC/Nextflow/Snakemake 运行时、Manubot、statcheck、Deepchecks、通用 Skill loader。",
    ]
    for index, gap in enumerate(gaps, 1):
        add(f"{index}. {gap}")
    add("")
    add("七、审计入口")
    add("=" * 72)
    add("产品能力：python scripts/audit_capability_registry.py --root .")
    add("上游采用：python scripts/audit_upstream_capabilities.py --repository-root .")
    add("权威来源：research_forge/capability_manifest.yaml 与 research_forge/upstream_capabilities.py")
    add("")
    add("结论：Research Forge 已经是一个‘有限实验类型下可闭环、全过程可审计’的平台；尚不是‘任意学科全自动完成且已被外部独立复现’的平台。")
    return "\n".join(lines) + "\n"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_first_row_as_header(table) -> None:
    """Mark the first row as a repeatable semantic table header."""
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def set_run_font(run, size: float = 10.5, color: str = INK, bold: bool = False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def style_paragraph(paragraph, after=6, line=1.25, color=INK, size=10.5) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bool(run.bold))


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def add_status_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 1800, 4860])
    mark_first_row_as_header(table)
    headers = ["能力", "状态", "说明"]
    for index, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], DEEP)
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    for name, state, note in rows:
        cells = table.add_row().cells
        for index, value in enumerate((name, state, note)):
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[index], LIGHT)
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            color = GREEN if index == 1 and "已达到" in value else (AMBER if index == 1 else INK)
            set_run_font(r, size=9.1, color=color, bold=index in (0, 1))
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
    set_table_geometry(table, [2700, 1800, 4860])
    doc.add_paragraph()


def build_docx() -> None:
    cap_audit = audit_capability_registry(ROOT)
    up_audit = audit_upstream_capabilities(ROOT)
    status_counts = Counter(item.status for item in UPSTREAM_CAPABILITIES)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading_tokens = {
        1: (16, DEEP, 18, 10),
        2: (13, MID, 14, 7),
        3: (11.5, DEEP, 10, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "RESEARCH FORGE  ·  CAPABILITY CATALOG"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(header, after=0, line=1.0, color=MUTED, size=8.5)
    footer = section.footer.paragraphs[0]
    footer.text = f"Research Forge · {DATE} · 当前能力以机器注册表和验收记录为准"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(footer, after=0, line=1.0, color=MUTED, size=8)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("CAPABILITY & UPSTREAM CATALOG"), size=9.5, color=MID, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    set_run_font(title.add_run("Research Forge 现有功能与 GitHub 能力总表"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("当前平台能力、上游项目作用、接入成熟度与未完成边界"), size=12, color=MUTED)

    summary = doc.add_table(rows=1, cols=3)
    summary.style = "Table Grid"
    set_table_geometry(summary, [3120, 3120, 3120])
    mark_first_row_as_header(summary)
    summary_values = [
        (f"{cap_audit.supported_count}/{cap_audit.manifest_count}", "产品能力声明有本地证据"),
        (f"{up_audit.supported_count}/{up_audit.record_count}", "上游采用状态通过定位审计"),
        ("C4", "当前最高真实案例等级；尚无 C5"),
    ]
    for cell, (value, label) in zip(summary.rows[0].cells, summary_values):
        set_cell_shading(cell, PALE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(value), size=18, color=DEEP, bold=True)
        p.add_run("\n")
        set_run_font(p.add_run(label), size=8.5, color=MUTED)
    doc.add_paragraph()

    add_heading(doc, "1. 结论摘要", 1)
    p = doc.add_paragraph("Research Forge 已形成可运行的四阶段研究工作台，并在流程控制、检索、证据约束写作、科研绘图、部分实验 Profile 和复现检查上取得真实验收。当前最强结论是：在已支持的计算实验类型中，平台可以把方向、契约、执行、证据、判定、修复和论文保持在同一条可审计链上。")
    style_paragraph(p)
    p = doc.add_paragraph("它尚不能被描述为“任意学科全自动科研系统”。本地 clean-room replay 已通过，但缺少独立操作者与外部签署收据，因此不达到 C5。")
    style_paragraph(p, color=AMBER)

    add_heading(doc, "2. 成熟度口径", 1)
    maturity = [
        ("C0", "概念或设计"), ("C1", "已有实现"), ("C2", "组件验证"),
        ("C3", "受控端到端"), ("C4", "真实案例验证"), ("C5", "外部独立信任域复现"),
    ]
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    set_table_geometry(table, [1560] * 6)
    mark_first_row_as_header(table)
    for cell, (code, label) in zip(table.rows[0].cells, maturity):
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(code), size=11, color=DEEP, bold=True)
        p.add_run("\n")
        set_run_font(p.add_run(label), size=8, color=MUTED)
    doc.add_paragraph()

    add_heading(doc, "3. 四阶段当前能力", 1)
    for title_text, flow, state in PHASES:
        add_heading(doc, title_text, 2)
        p = doc.add_paragraph(flow)
        style_paragraph(p)
        p = doc.add_paragraph("当前状态：" + state)
        style_paragraph(p, color=DEEP, size=9.5)

    add_heading(doc, "4. 核心功能清单", 1)
    add_status_table(doc, [(name, status, ability + " 边界：" + boundary) for name, status, ability, boundary in CORE_CAPABILITIES])

    add_heading(doc, "5. GitHub / Skill / 标准采用总览", 1)
    p = doc.add_paragraph("采用状态描述的是 Research Forge 与上游能力的真实关系，不是宣传性功能计数。尤其要区分：架构参考、Codex 会话可用 Skill、产品适配器和真实运行时接入。")
    style_paragraph(p)
    overview_rows = []
    for status in AdoptionStatus:
        names = "、".join(item.name for item in UPSTREAM_CAPABILITIES if item.status is status)
        overview_rows.append((STATUS_CN[status], str(status_counts.get(status, 0)), names or "—"))
    add_status_table(doc, [(name, count + " 项", names) for name, count, names in overview_rows])

    add_heading(doc, "6. 逐项上游能力与接入情况", 1)
    for status in AdoptionStatus:
        records = [item for item in UPSTREAM_CAPABILITIES if item.status is status]
        if not records:
            continue
        add_heading(doc, f"{STATUS_CN[status]}（{len(records)} 项）", 2)
        for item in records:
            add_heading(doc, item.name, 3)
            source = doc.add_paragraph()
            set_run_font(source.add_run("来源："), size=9, color=MUTED, bold=True)
            set_run_font(source.add_run(item.source_url), size=9, color=MID)
            source.paragraph_format.space_after = Pt(3)
            detail = doc.add_paragraph()
            set_run_font(detail.add_run("上游能力："), size=9.5, color=DEEP, bold=True)
            set_run_font(detail.add_run(compact("；".join(item.upstream_capabilities), 520)), size=9.5)
            detail.paragraph_format.space_after = Pt(3)
            adopted = doc.add_paragraph()
            set_run_font(adopted.add_run("平台采用："), size=9.5, color=DEEP, bold=True)
            learned = "；".join(item.learned_capabilities) or "未转化为平台能力"
            set_run_font(adopted.add_run(compact(learned, 420)), size=9.5)
            adopted.paragraph_format.space_after = Pt(3)
            claim = doc.add_paragraph()
            set_run_font(claim.add_run("允许声明："), size=9.5, color=GREEN, bold=True)
            set_run_font(claim.add_run(item.permitted_claim), size=9.5)
            claim.paragraph_format.space_after = Pt(3)
            bound = doc.add_paragraph()
            set_run_font(bound.add_run("限制："), size=9.5, color=AMBER, bold=True)
            set_run_font(bound.add_run(compact("；".join(item.missing_or_bounded), 520)), size=9.5)
            bound.paragraph_format.space_after = Pt(6)

    add_heading(doc, "7. 最重要的未完成项", 1)
    gaps = [
        ("外部独立复现（C5）", "需要独立操作者/信任域只使用冻结包复现，并生成可验证的外部收据。"),
        ("通用 Contract Compiler", "继续扩展实验 Profile、可执行指标、目标规则、采样框与资源授权。"),
        ("Provenance Run RO-Crate", "Workflow Run 0.5 已验收；Provenance Run Profile 仍待完成。"),
        ("Nuwa 与 next-ai-draw-io", "分别补真实模型/负责人裁决 UI，以及真实 MCP 对话编辑端到端验收。"),
        ("规模化全文证据 benchmark", "补多文献冲突、攻击、弃权/引用准确率和 P95/P99。"),
        ("未接入的外部运行时", "OSF 在线注册、MLflow、DVC/Nextflow/Snakemake、Manubot、statcheck、Deepchecks 和通用 Skill loader。"),
    ]
    add_status_table(doc, [(name, "未完成", note) for name, note in gaps])

    add_heading(doc, "8. 审计入口", 1)
    for command, note in [
        ("python scripts/audit_capability_registry.py --root .", "审计产品能力声明与证据等级。"),
        ("python scripts/audit_upstream_capabilities.py --repository-root .", "审计 GitHub/Skill 采用状态与证据定位。"),
        ("research_forge/capability_manifest.yaml", "产品能力事实来源。"),
        ("research_forge/upstream_capabilities.py", "上游项目采用状态事实来源。"),
    ]:
        p = doc.add_paragraph()
        set_run_font(p.add_run(command), size=9.5, color=DEEP, bold=True)
        p.add_run("\n")
        set_run_font(p.add_run(note), size=9, color=MUTED)
        p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph("最终判断：Research Forge 已经是一个“有限实验类型下可闭环、全过程可审计”的平台；尚不是“任意学科全自动完成且已被外部独立复现”的平台。")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(), size=10.5)
    for run in p.runs:
        set_run_font(run, size=10.5, color=DEEP, bold=True)

    doc.core_properties.title = "Research Forge 现有功能与 GitHub 能力总表"
    doc.core_properties.subject = "平台能力、上游项目采用状态与证据边界"
    doc.core_properties.author = "Research Forge"
    doc.core_properties.keywords = "Research Forge, GitHub, capability, audit"
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main() -> None:
    TXT_PATH.parent.mkdir(parents=True, exist_ok=True)
    TXT_PATH.write_text(build_txt(), encoding="utf-8-sig")
    build_docx()
    print(TXT_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
