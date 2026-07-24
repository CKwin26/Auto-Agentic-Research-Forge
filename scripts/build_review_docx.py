from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "review" / "论文匿名评审表.docx"


def set_run_font(run, size=8.5, bold=False):
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_cell(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    set_run_font(paragraph.add_run(text), bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(rows[0]):
        set_cell(table.rows[0].cells[index], text, True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF7")
        table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shading)
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            set_cell(cells[index], text)
    return table


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    document.styles["Normal"].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("论文匿名评审表"), size=16, bold=True)
    intro = document.add_paragraph(
        "请仅依据收到的匿名论文评审，不搜索作者或项目。预计用时：10–15 分钟。"
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("一、基本信息", level=2)
    add_table(
        document,
        [
            ["项目", "填写"],
            ["评审编号", ""],
            ["相关领域", ""],
            ["熟悉程度", "□ 5 很熟悉　□ 4　□ 3　□ 2　□ 1 不熟悉"],
            ["利益冲突", "□ 无　□ 有（请说明）："],
        ],
    )

    document.add_heading("二、快速评分", level=2)
    document.add_paragraph(
        "评分：1 = 很差，3 = 基本合格，5 = 很好；无法判断请填“不确定”。"
    )
    dimensions = [
        "研究问题重要性",
        "创新性与贡献清晰度",
        "方法与实验设计严谨性",
        "结果和统计报告可信度",
        "论文主张与证据一致性",
        "局限与结论边界是否诚实",
        "写作、结构与图表清晰度",
    ]
    add_table(
        document,
        [["维度", "1–5", "一句话理由（请标页码）"]]
        + [[item, "", ""] for item in dimensions],
    )

    document.add_heading("三、关键审计", level=2)
    document.add_paragraph(
        "请勾选：A = 支持，B = 不支持，C = 论文信息不足、无法判断。"
    )
    claims = [
        "同骨干、共享上游工件的配对设计能够隔离下游门控影响",
        "自动结果只能视为代理指标，不能证明门控具有确认性收益",
        "五条告警的原因被合理定位为：3 条上下文打包缺失、2 条判定优先级错误",
        "恢复上下文后五条均获支持，只能用于诊断，不能覆盖原始盲审结果",
        "论文没有把五案例回放夸大为前瞻性泛化证据",
        "核心贡献是“可审计的评估与故障定位框架”，而非“门控效果已获证明”",
    ]
    add_table(
        document,
        [["待审计陈述", "A", "B", "C", "依据或缺失信息（页码）"]]
        + [[item, "□", "□", "□", ""] for item in claims],
    )

    document.add_heading("四、总体意见", level=2)
    prompts = [
        ("最强优点（1–2 条）", 2),
        ("必须修改的问题（最多 3 条；写清位置、影响和建议）", 3),
        ("你认为最强的反对理由", 1),
        ("需要作者回答的一个问题", 1),
    ]
    for label, line_count in prompts:
        paragraph = document.add_paragraph()
        set_run_font(paragraph.add_run(label + "："), bold=True)
        for index in range(line_count):
            prefix = f"{index + 1}. " if line_count > 1 else ""
            document.add_paragraph(prefix + "_" * 78)

    document.add_heading("五、最终结论", level=2)
    add_table(
        document,
        [
            ["项目", "选择"],
            ["投稿建议", "□ 接收　□ 小修　□ 大修　□ 拒稿"],
            [
                "更适合的发表形式",
                "□ 正式会议/期刊　□ Workshop/小型会议　□ 预印本后补实验　□ 暂不投稿",
            ],
            ["是否值得继续完善", "□ 是　□ 否　□ 不确定"],
        ],
    )
    paragraph = document.add_paragraph()
    set_run_font(paragraph.add_run("结论理由（不超过 100 字）："), bold=True)
    document.add_paragraph("_" * 90)
    document.add_paragraph("_" * 90)
    document.add_paragraph("是否同意匿名引用本评审意见：　□ 同意　□ 不同意")

    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(1)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
